import datetime
import io

import openpyxl
from django.contrib.auth.models import Group, User
from django.test import TestCase

from .models import Report, Section, VehicleRecord


class AnalyticsExportTests(TestCase):
    def setUp(self):
        analyst_group = Group.objects.create(name='Аналитика')
        self.analyst = User.objects.create_user(username='analyst', password='pass')
        self.analyst.groups.add(analyst_group)
        self.section = Section.objects.create(name='Карьер')
        self.report = Report.objects.create(
            name='Проверочный отчёт',
            section=self.section,
            daily_norm_sec=36000,
            bulldozer_norm_sec=10800,
            excavator_norm_sec=13200,
            dumptruck_norm_sec=10800,
        )
        self.record = VehicleRecord.objects.create(
            report=self.report,
            row_number=1,
            name='Volvo A40F №352',
            group='Самосвалы',
            date='08.09',
            record_date=datetime.date.today(),
            shift=1,
            engine_time_sec=3600,
            fuel_actual=120,
            fuel_norm=30,
            fuel_efficiency=0.8,
            mileage=54.5,
            refueling=25,
        )
        self.client.login(username='analyst', password='pass')

    def test_analytics_defaults_to_last_30_days(self):
        response = self.client.get('/analytics/')
        today = datetime.date.today()

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['date_from'], (today - datetime.timedelta(days=30)).isoformat())
        self.assertEqual(response.context['date_to'], today.isoformat())

    def test_analytics_subsections_have_separate_pages(self):
        general = self.client.get('/analytics/')
        fuel = self.client.get('/analytics/fuel/')
        mileage = self.client.get('/analytics/mileage/')

        self.assertEqual(general.context['analytics_page'], 'general')
        self.assertEqual(fuel.context['analytics_page'], 'fuel')
        self.assertEqual(mileage.context['analytics_page'], 'mileage')
        self.assertContains(fuel, 'Аналитика топлива')
        self.assertNotContains(fuel, 'Общая аналитика эффективности')
        self.assertContains(mileage, 'Аналитика пробега')
        self.assertNotContains(mileage, 'Общая аналитика эффективности')

    def test_vehicle_search_matches_name_and_board_number(self):
        by_name = self.client.get('/analytics/', {
            'date_from': '2000-01-01',
            'date_to': '2100-01-01',
            'vehicle_search': 'volvo',
        })
        by_number = self.client.get('/analytics/', {
            'date_from': '2000-01-01',
            'date_to': '2100-01-01',
            'vehicle_search': '352',
        })

        self.assertEqual(by_name.context['total_count'], 1)
        self.assertEqual(by_number.context['total_count'], 1)

    def test_vehicle_search_is_unicode_case_insensitive(self):
        self.record.name = 'БелАЗ №352'
        self.record.save(update_fields=['name'])

        response = self.client.get('/analytics/', {
            'date_from': '2000-01-01',
            'date_to': '2100-01-01',
            'vehicle_search': 'белаз',
        })

        self.assertEqual(response.context['total_count'], 1)

    def test_export_uses_allowlisted_selected_columns(self):
        response = self.client.get('/analytics/export/', {
            'dataset': 'fuel',
            'date_from': '2000-01-01',
            'date_to': '2100-01-01',
            'columns': ['vehicle', 'fuel_actual', 'password'],
        })

        self.assertEqual(response.status_code, 200)
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        headers = [cell.value for cell in workbook.active[1]]
        self.assertEqual(headers, ['Техника', 'Фактический расход, л'])

    def test_export_supports_legacy_vehicle_filter(self):
        VehicleRecord.objects.create(
            report=self.report,
            row_number=2,
            name='CAT 777 №100',
            group='Самосвалы',
            date='08.09',
            record_date=datetime.date.today(),
            shift=1,
            engine_time_sec=3600,
            fuel_actual=80,
            fuel_norm=30,
        )

        response = self.client.get('/analytics/export/', {
            'dataset': 'fuel',
            'vehicle': 'volvo',
            'columns': ['vehicle'],
        })
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        data_rows = list(workbook.active.iter_rows(min_row=2, values_only=True))

        self.assertEqual(data_rows, [('Volvo A40F №352',)])

    def test_mileage_export_excludes_missing_values(self):
        VehicleRecord.objects.create(
            report=self.report,
            row_number=2,
            name='Volvo A40F №353',
            group='Самосвалы',
            date='08.09',
            record_date=datetime.date.today(),
            shift=1,
            engine_time_sec=3600,
            fuel_norm=30,
            mileage=None,
        )

        response = self.client.get('/analytics/export/', {
            'dataset': 'mileage',
            'date_from': '2000-01-01',
            'date_to': '2100-01-01',
            'columns': ['vehicle', 'mileage'],
        })
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        data_rows = list(workbook.active.iter_rows(min_row=2, values_only=True))

        self.assertEqual(data_rows, [('Volvo A40F №352', 54.5)])

    def test_direct_export_defaults_to_last_30_days(self):
        VehicleRecord.objects.create(
            report=self.report,
            row_number=3,
            name='Старая техника №1',
            group='Самосвалы',
            date='01.01',
            record_date=datetime.date.today() - datetime.timedelta(days=90),
            shift=1,
            engine_time_sec=3600,
            fuel_actual=10,
            fuel_norm=30,
        )

        response = self.client.get('/analytics/export/', {
            'dataset': 'fuel',
            'columns': ['vehicle'],
        })
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        data_rows = list(workbook.active.iter_rows(min_row=2, values_only=True))

        self.assertEqual(data_rows, [('Volvo A40F №352',)])

    def test_export_escapes_formula_leading_text(self):
        self.record.name = '=HYPERLINK(\"https://example.test\")'
        self.record.save(update_fields=['name'])

        response = self.client.get('/analytics/export/', {
            'dataset': 'fuel',
            'columns': ['vehicle'],
        })
        workbook = openpyxl.load_workbook(io.BytesIO(response.content), data_only=False)
        value = workbook.active.cell(row=2, column=1).value

        self.assertEqual(value, '\'=HYPERLINK(\"https://example.test\")')

    def test_summary_export_combines_shifts_per_vehicle_and_day(self):
        VehicleRecord.objects.create(
            report=self.report,
            row_number=2,
            name=self.record.name,
            group=self.record.group,
            date=self.record.date,
            record_date=self.record.record_date,
            shift=2,
            engine_time_sec=7200,
            fuel_actual=180,
            fuel_norm=60,
            mileage=45.5,
            refueling=None,
        )

        response = self.client.get('/analytics/export/', {
            'dataset': 'fuel',
            'export_mode': 'summary',
            'columns': [
                'date', 'shift', 'vehicle', 'engine_hours',
                'fuel_actual', 'fuel_norm', 'fuel_efficiency', 'refueling',
            ],
        })
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        rows = list(workbook.active.iter_rows(values_only=True))

        self.assertEqual(
            rows[0],
            (
                'Дата', 'Техника', 'Время работы, ч',
                'Фактический расход, л', 'Норма расхода, л/ч',
                'Расход к норме, %', 'Заправка, л',
            ),
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][1], 'Volvo A40F №352')
        self.assertEqual(rows[1][2], 3)
        self.assertEqual(rows[1][3], 300)
        self.assertEqual(rows[1][4], 50)
        self.assertEqual(rows[1][5], 200)
        self.assertEqual(rows[1][6], 25)

    def test_mileage_summary_preserves_missing_values_and_sums_shifts(self):
        VehicleRecord.objects.create(
            report=self.report,
            row_number=2,
            name=self.record.name,
            group=self.record.group,
            date=self.record.date,
            record_date=self.record.record_date,
            shift=2,
            engine_time_sec=3600,
            fuel_norm=30,
            mileage=10.5,
        )

        response = self.client.get('/analytics/export/', {
            'dataset': 'mileage',
            'export_mode': 'summary',
            'columns': ['date', 'shift', 'vehicle', 'mileage'],
        })
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        rows = list(workbook.active.iter_rows(values_only=True))

        self.assertEqual(rows[0], ('Дата', 'Техника', 'Пробег, км'))
        self.assertEqual(rows[1][2], 65)

    def test_summary_keeps_legacy_same_day_from_different_years_separate(self):
        self.record.record_date = None
        self.record.save(update_fields=['record_date'])
        old_report = Report.objects.create(
            name='Старый отчёт',
            section=self.section,
            year=self.report.year - 1,
            daily_norm_sec=36000,
            bulldozer_norm_sec=10800,
            excavator_norm_sec=13200,
            dumptruck_norm_sec=10800,
        )
        VehicleRecord.objects.create(
            report=old_report,
            row_number=1,
            name=self.record.name,
            group=self.record.group,
            date=self.record.date,
            record_date=None,
            shift=2,
            engine_time_sec=3600,
            fuel_actual=50,
            fuel_norm=30,
        )

        response = self.client.get('/analytics/export/', {
            'dataset': 'fuel',
            'export_mode': 'summary',
            'date_from': 'invalid',
            'date_to': 'invalid',
            'columns': ['date', 'vehicle', 'fuel_actual'],
        })
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        rows = list(workbook.active.iter_rows(min_row=2, values_only=True))

        self.assertEqual(len(rows), 2)